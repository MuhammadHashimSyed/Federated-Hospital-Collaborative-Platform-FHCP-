"""
FYP Report Generator — LaTeX-style professional DOCX
Table design matches FYP_Report_111.docx exactly:
  - Booktabs style: thick top rule (sz=12), thin header-separator rule (sz=4),
    thin bottom rule (sz=6), NO fills, NO vertical borders
  - Palatino Linotype body/table font (matches original)
  - Bold centered headers, left-aligned data cells
  - No background shading anywhere in tables
"""

import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CHAPTERS_DIR = r"c:\Users\arbab\Desktop\myfrontend\report_chapters"
OUTPUT_FILE  = r"c:\Users\arbab\Desktop\myfrontend\fyp_report_ext.docx"

BODY_FONT  = "Palatino Linotype"   # matches original report
HEAD_FONT  = "Times New Roman"     # headings / cover page
CODE_FONT  = "Courier New"

# ─────────────────────────────────────────────────────────────
# LOW-LEVEL HELPERS
# ─────────────────────────────────────────────────────────────

def set_page_margins(section, top=2.54, bottom=2.54, left=3.0, right=2.54):
    section.top_margin    = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin   = Cm(left)
    section.right_margin  = Cm(right)

def set_para_spacing(para, before=0, after=6, line=None):
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after  = Pt(after)
    if line:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing      = line

def add_run(para, text, bold=False, italic=False, size=12,
            font=None, color=None):
    font = font or BODY_FONT
    run = para.add_run(text)
    run.bold        = bold
    run.italic      = italic
    run.font.name   = font
    run.font.size   = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run

def add_border_bottom(para, color="2C3E50", sz="6"):
    """Paragraph bottom border — used as a decorative horizontal rule."""
    pPr = para._p.get_or_add_pPr()
    # Check if pBdr exists
    pBdr = pPr.find(qn('w:pBdr'))
    if pBdr is None:
        pBdr = OxmlElement('w:pBdr')
        pPr.append(pBdr)
    
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),   'single')
    bottom.set(qn('w:sz'),    sz)
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)

def add_horizontal_rule(doc):
    para = doc.add_paragraph()
    set_para_spacing(para, before=4, after=4)
    add_border_bottom(para)

def add_page_break(doc):
    doc.add_page_break()

# ─────────────────────────────────────────────────────────────
# HEADINGS
# ─────────────────────────────────────────────────────────────

def make_heading(doc, text, level=1):
    if level == 0:          # Chapter title
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_para_spacing(para, before=24, after=6)
        add_run(para, text.upper(), bold=True, size=14, font=HEAD_FONT)
        rule = doc.add_paragraph()
        set_para_spacing(rule, before=0, after=12)
        add_border_bottom(rule, color="000000", sz="8")

    elif level == 1:        # Section  (##)
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_para_spacing(para, before=14, after=4)
        add_run(para, text, bold=True, size=13, font=HEAD_FONT)

    elif level == 2:        # Subsection  (###)
        para = doc.add_paragraph()
        set_para_spacing(para, before=10, after=3)
        add_run(para, text, bold=True, size=12, font=HEAD_FONT)

    elif level == 3:        # Sub-subsection  (####)
        para = doc.add_paragraph()
        set_para_spacing(para, before=8, after=2)
        add_run(para, text, bold=True, italic=True, size=11, font=HEAD_FONT)

# ─────────────────────────────────────────────────────────────
# TABLE — EXACT BOOKTABS MATCH (FYP_Report_111 style)
# ─────────────────────────────────────────────────────────────

def _get_or_add_tcPr(cell):
    tc = cell._tc
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr')
        tc.insert(0, tcPr)
    return tcPr

def _get_or_add_tblPr(table):
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    return tblPr

def _set_cell_border(cell, top=None, bottom=None, left=None, right=None):
    """Set individual cell borders."""
    tcPr = _get_or_add_tcPr(cell)
    tcBdr = tcPr.find(qn('w:tcBorders'))
    if tcBdr is None:
        tcBdr = OxmlElement('w:tcBorders')
        tcPr.append(tcBdr)
    
    for edge, spec in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if spec is None: continue
        # Remove existing if any
        existing = tcBdr.find(qn(f'w:{edge}'))
        if existing is not None:
            tcBdr.remove(existing)
            
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'),   spec.get('val',   'none'))
        el.set(qn('w:sz'),    spec.get('sz',    '0'))
        el.set(qn('w:color'), spec.get('color', 'auto'))
        el.set(qn('w:space'), '0')
        tcBdr.append(el)

def _clear_table_outer_borders(table):
    """Remove all outer and inner table-level borders."""
    tblPr = _get_or_add_tblPr(table)
    tblBdr = tblPr.find(qn('w:tblBorders'))
    if tblBdr is None:
        tblBdr = OxmlElement('w:tblBorders')
        tblPr.append(tblBdr)
    else:
        # Clear existing
        for child in tblBdr.getchildren():
            tblBdr.remove(child)

    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'),   'none')
        el.set(qn('w:sz'),    '0')
        el.set(qn('w:color'), 'auto')
        tblBdr.append(el)

NONE_BORDER   = {'val': 'none',   'sz': '0',  'color': 'FFFFFF'}
THICK_BORDER  = {'val': 'single', 'sz': '12', 'color': '000000'}  # top of header
THIN_BORDER   = {'val': 'single', 'sz': '4',  'color': '000000'}  # below header
CLOSE_BORDER  = {'val': 'single', 'sz': '6',  'color': '000000'}  # table bottom

def add_table_from_md(doc, rows):
    """
    Render a markdown table with exact FYP_Report_111 booktabs styling:
      - No fills, no shading anywhere
      - Header: thick top rule (sz=12) + thin bottom rule (sz=4), bold, centered
      - Data rows: no borders (open rows)
      - Last data row: thin bottom closing rule (sz=6)
      - Font: Palatino Linotype 10pt (matches original)
    """
    if len(rows) < 2:
        return

    header = rows[0]
    data   = rows[2:] if len(rows) > 2 else []
    ncols  = len(header)

    table = doc.add_table(rows=1 + len(data), cols=ncols)
    # Using 'Normal' style often helps avoiding unexpected grid lines from named styles
    table.style = None 
    _clear_table_outer_borders(table)

    # ── Header row ────────────────────────────────
    hdr_row = table.rows[0]
    for col_i, txt in enumerate(header):
        cell = hdr_row.cells[col_i]
        cell.text = ""
        p  = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para_spacing(p, before=3, after=3)
        run = p.add_run(txt.strip())
        run.bold          = True
        run.font.name     = BODY_FONT
        run.font.size     = Pt(10)
        _set_cell_border(cell,
                         top    = THICK_BORDER,
                         bottom = THIN_BORDER,
                         left   = NONE_BORDER,
                         right  = NONE_BORDER)

    # ── Data rows ─────────────────────────────────
    for row_i, row_data in enumerate(data):
        tr       = table.rows[row_i + 1]
        is_last  = (row_i == len(data) - 1)
        bot_spec = CLOSE_BORDER if is_last else NONE_BORDER
        for col_i, txt in enumerate(row_data):
            if col_i >= ncols:
                break
            cell = tr.cells[col_i]
            cell.text = ""
            p  = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_para_spacing(p, before=2, after=2)
            run = p.add_run(txt.strip())
            run.font.name = BODY_FONT
            run.font.size = Pt(10)
            _set_cell_border(cell,
                             top    = NONE_BORDER,
                             bottom = bot_spec,
                             left   = NONE_BORDER,
                             right  = NONE_BORDER)

    # Space after table
    sp = doc.add_paragraph()
    set_para_spacing(sp, before=0, after=8)

# ─────────────────────────────────────────────────────────────
# CODE BLOCK
# ─────────────────────────────────────────────────────────────

def add_code_block(doc, text):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para_spacing(para, before=4, after=6)
    pf = para.paragraph_format
    pf.left_indent  = Cm(0.6)
    pf.right_indent = Cm(0.6)
    
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  'F2F3F4')
    pPr.append(shd)
    
    run = para.add_run(text)
    run.font.name = CODE_FONT
    run.font.size = Pt(9)

# ─────────────────────────────────────────────────────────────
# INLINE FORMATTING
# ─────────────────────────────────────────────────────────────

def inline_format(para, text, size=11):
    pattern = re.compile(r'(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            r = para.add_run(text[pos:m.start()])
            r.font.name = BODY_FONT
            r.font.size = Pt(size)
        full = m.group(0)
        if full.startswith('**'):
            r = para.add_run(m.group(2))
            r.bold = True
            r.font.name = BODY_FONT
            r.font.size = Pt(size)
        elif full.startswith('*'):
            r = para.add_run(m.group(3))
            r.italic = True
            r.font.name = BODY_FONT
            r.font.size = Pt(size)
        else:
            r = para.add_run(m.group(4))
            r.font.name = CODE_FONT
            r.font.size = Pt(size - 1)
        pos = m.end()
    if pos < len(text):
        r = para.add_run(text[pos:])
        r.font.name = BODY_FONT
        r.font.size = Pt(size)

# ─────────────────────────────────────────────────────────────
# MARKDOWN PARSER
# ─────────────────────────────────────────────────────────────

def parse_md_to_docx(doc, md_text):
    lines        = md_text.split('\n')
    i            = 0
    in_code      = False
    code_lines   = []
    table_rows   = []
    in_table     = False

    while i < len(lines):
        line    = lines[i]
        stripped = line.strip()

        if stripped.startswith('```'):
            if not in_code:
                in_code    = True
                code_lines = []
            else:
                in_code = False
                if code_lines:
                    add_code_block(doc, '\n'.join(code_lines))
                code_lines = []
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if stripped.startswith('|'):
            parts = [p.strip() for p in stripped.split('|')]
            if len(parts) > 2:
                parts = parts[1:-1]
            table_rows.append(parts)
            in_table = True
            i += 1
            continue
        else:
            if in_table and table_rows:
                add_table_from_md(doc, table_rows)
                table_rows = []
                in_table   = False

        if not stripped:
            i += 1
            continue

        if re.match(r'^-{3,}$', stripped):
            add_horizontal_rule(doc)
            i += 1
            continue

        if line.startswith('# ') and not line.startswith('## '):
            make_heading(doc, line[2:].strip(), level=0)
            i += 1
            continue

        if line.startswith('## '):
            make_heading(doc, line[3:].strip(), level=1)
            i += 1
            continue

        if line.startswith('### '):
            make_heading(doc, line[4:].strip(), level=2)
            i += 1
            continue

        if line.startswith('#### '):
            make_heading(doc, line[5:].strip(), level=3)
            i += 1
            continue

        if re.match(r'^[-*] ', stripped):
            text = stripped[2:]
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_para_spacing(para, before=1, after=2)
            para.paragraph_format.left_indent = Cm(0.8)
            r0 = para.add_run('• ')
            r0.font.name = BODY_FONT
            r0.font.size = Pt(11)
            r0.bold      = True
            inline_format(para, text)
            i += 1
            continue

        if re.match(r'^\d+\. ', stripped):
            text = re.sub(r'^\d+\. ', '', stripped)
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_para_spacing(para, before=1, after=2)
            para.paragraph_format.left_indent = Cm(0.8)
            inline_format(para, text)
            i += 1
            continue

        if stripped.startswith('*End of') or stripped.startswith('*end of'):
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_para_spacing(para, before=12, after=0)
            r = para.add_run(stripped.strip('*'))
            r.italic      = True
            r.font.name   = BODY_FONT
            r.font.size   = Pt(10)
            r.font.color.rgb = RGBColor(120, 120, 120)
            i += 1
            continue

        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_para_spacing(para, before=0, after=6, line=1.15)
        inline_format(para, stripped)
        i += 1

    if in_table and table_rows:
        add_table_from_md(doc, table_rows)

# ─────────────────────────────────────────────────────────────
# COVER PAGE
# ─────────────────────────────────────────────────────────────

def add_cover_page(doc):
    section = doc.sections[0]
    set_page_margins(section, top=3.0, bottom=3.0, left=3.0, right=2.54)

    def cp(text, size, bold=False, color=(0,0,0), after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para_spacing(p, before=0, after=after)
        r = p.add_run(text)
        r.font.name  = HEAD_FONT
        r.font.size  = Pt(size)
        r.bold       = bold
        r.font.color.rgb = RGBColor(*color)
        return p

    cp("BAHRIA UNIVERSITY ISLAMABAD", 16, bold=True, color=(28,40,51), after=4)
    cp("Department of Computer Science", 12, color=(52,73,94), after=8)
    add_horizontal_rule(doc)

    for _ in range(3):
        doc.add_paragraph()

    cp("FINAL YEAR PROJECT REPORT", 13, bold=True, color=(52,73,94), after=20)

    for _ in range(2):
        doc.add_paragraph()

    cp("FEDERATED HOSPITAL COLLABORATIVE PLATFORM:", 18, bold=True, color=(28,40,51), after=8)
    cp("A WEB-BASED IMPLEMENTATION FOR PRIVACY-PRESERVING", 14, bold=True, color=(28,40,51), after=6)
    cp("CHEST X-RAY DIAGNOSIS", 14, bold=True, color=(28,40,51), after=30)

    add_horizontal_rule(doc)
    for _ in range(3):
        doc.add_paragraph()

    for label, val in [
        ("Submitted By:", ""),
        ("", "Arbab Hussain           (01-134222-173)"),
        ("", "Muhammad Hashim Syed    (01-134222-095)"),
        ("", ""),
        ("Supervisor:", "Dr. Sumaira Kausar"),
        ("", ""),
        ("Session:", "April 2026"),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para_spacing(p, before=0, after=4)
        if label:
            r = p.add_run(label + "  ")
            r.font.name = HEAD_FONT
            r.font.size = Pt(12)
            r.bold = True
        if val:
            r2 = p.add_run(val)
            r2.font.name = HEAD_FONT
            r2.font.size = Pt(12)

    add_page_break(doc)

# ─────────────────────────────────────────────────────────────
# HEADER / FOOTER
# ─────────────────────────────────────────────────────────────

def add_header_footer(doc,
                      title="Federated Hospital Collaborative Platform — FYP Report"):
    for section in doc.sections[1:]:
        hdr = section.header
        hp  = hdr.paragraphs[0] if hdr.paragraphs else hdr.add_paragraph()
        hp.clear()
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run(title)
        r.font.name = HEAD_FONT
        r.font.size = Pt(9)
        r.italic    = True
        r.font.color.rgb = RGBColor(130, 130, 130)
        
        pPr    = hp._p.get_or_add_pPr()
        pBdr   = pPr.find(qn('w:pBdr'))
        if pBdr is None:
            pBdr = OxmlElement('w:pBdr')
            pPr.append(pBdr)
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'),   'single')
        bottom.set(qn('w:sz'),    '4')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), 'AAAAAA')
        pBdr.append(bottom)

        ftr = section.footer
        fp  = ftr.paragraphs[0] if ftr.paragraphs else ftr.add_paragraph()
        fp.clear()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = fp.add_run()
        r2.font.name = HEAD_FONT
        r2.font.size = Pt(9)
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.text = 'PAGE'
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        r2._r.append(fldChar1)
        r2._r.append(instrText)
        r2._r.append(fldChar2)

# ─────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────

def build_report():
    if not os.path.exists(CHAPTERS_DIR):
        print(f"Error: Chapters directory {CHAPTERS_DIR} not found.")
        return

    doc   = Document()
    style = doc.styles['Normal']
    style.font.name = BODY_FONT
    style.font.size = Pt(11)

    add_cover_page(doc)

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_page_margins(doc.sections[-1], top=2.54, bottom=2.54, left=3.0, right=2.54)

    chapter_files = [
        "Chapter_0_FrontMatter.md",
        "Chapter_1_Introduction.md",
        "Chapter_2_Literature_Review.md",
        "Chapter_3_Requirements.md",
        "Chapter_4_System_Design.md",
        "Chapter_5_Implementation.md",
        "Chapter_6_Testing_Evaluation.md",
        "Chapter_7_Conclusion_References.md",
    ]

    for idx, fname in enumerate(chapter_files):
        fpath = os.path.join(CHAPTERS_DIR, fname)
        if not os.path.exists(fpath):
            print(f"  MISSING: {fname}")
            continue
        with open(fpath, 'r', encoding='utf-8') as f:
            md_text = f.read()
        print(f"  Processing {fname} ...")
        parse_md_to_docx(doc, md_text)
        if idx < len(chapter_files) - 1:
            doc.add_section(WD_SECTION.NEW_PAGE)
            set_page_margins(doc.sections[-1], top=2.54, bottom=2.54, left=3.0, right=2.54)

    add_header_footer(doc)
    doc.save(OUTPUT_FILE)
    print(f"Report saved to: {OUTPUT_FILE}")

if __name__ == "__main__":
    print("Building FYP Report DOCX...")
    build_report()
