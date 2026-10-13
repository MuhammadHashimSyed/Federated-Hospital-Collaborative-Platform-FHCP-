from docx import Document
from docx.oxml.ns import qn

doc = Document(r'c:\Users\arbab\Desktop\myfrontend\FYP_Report_111.docx')

print(f'Total tables in document: {len(doc.tables)}')
print()

for t_idx, table in enumerate(doc.tables[:6]):
    print(f'=== TABLE {t_idx+1} ===')
    print(f'  Style name : {table.style.name}')
    print(f'  Rows: {len(table.rows)}, Cols: {len(table.columns)}')

    # Table-level XML
    tbl_pr = table._tbl.find(qn('w:tblPr'))
    if tbl_pr is not None:
        tbl_borders = tbl_pr.find(qn('w:tblBorders'))
        if tbl_borders is not None:
            for border in tbl_borders:
                tag = border.tag.split('}')[-1]
                val   = border.get(qn('w:val'), '')
                color = border.get(qn('w:color'), '')
                sz    = border.get(qn('w:sz'), '')
                print(f'    TblBorder [{tag}]: val={val}, color=#{color}, sz={sz}')

        shd = tbl_pr.find(qn('w:shd'))
        if shd is not None:
            print(f'    Table-level shade fill: #{shd.get(qn("w:fill"), "")}')

        look = tbl_pr.find(qn('w:tblLook'))
        if look is not None:
            print(f'    tblLook val: {look.get(qn("w:val"), "")}')

    # Header row
    if table.rows:
        row0 = table.rows[0]
        print(f'  -- Header row (row 0) --')
        for c_idx, cell in enumerate(row0.cells[:4]):
            tc_pr = cell._tc.find(qn('w:tcPr'))
            cell_info = []
            if tc_pr is not None:
                shd = tc_pr.find(qn('w:shd'))
                if shd is not None:
                    fill  = shd.get(qn('w:fill'), 'none')
                    theme = shd.get(qn('w:themeFill'), '')
                    cell_info.append(f'fill=#{fill} theme={theme}')
                borders = tc_pr.find(qn('w:tcBorders'))
                if borders is not None:
                    for b in borders:
                        tag   = b.tag.split('}')[-1]
                        val   = b.get(qn('w:val'), '')
                        color = b.get(qn('w:color'), '')
                        sz    = b.get(qn('w:sz'), '')
                        cell_info.append(f'border[{tag}]: val={val} color=#{color} sz={sz}')
            for para in cell.paragraphs:
                for run in para.runs:
                    fn = run.font.name or 'inherited'
                    fs = run.font.size
                    fb = run.bold
                    try:
                        fc = str(run.font.color.rgb) if run.font.color and run.font.color.type else 'inherited'
                    except Exception:
                        fc = 'inherited'
                    cell_info.append(f'font: {fn}, {fs}, bold={fb}, color={fc}')
                    break
                break
            print(f'    cell[{c_idx}]: {" | ".join(cell_info)}')

    # Data rows
    for r_idx in range(1, min(3, len(table.rows))):
        row = table.rows[r_idx]
        print(f'  -- Data row {r_idx} --')
        for c_idx, cell in enumerate(row.cells[:2]):
            tc_pr = cell._tc.find(qn('w:tcPr'))
            cell_info = []
            if tc_pr is not None:
                shd = tc_pr.find(qn('w:shd'))
                if shd is not None:
                    fill  = shd.get(qn('w:fill'), 'none')
                    theme = shd.get(qn('w:themeFill'), '')
                    cell_info.append(f'fill=#{fill} theme={theme}')
            for para in cell.paragraphs:
                for run in para.runs:
                    fn = run.font.name or 'inherited'
                    fs = run.font.size
                    try:
                        fc = str(run.font.color.rgb) if run.font.color and run.font.color.type else 'inherited'
                    except Exception:
                        fc = 'inherited'
                    cell_info.append(f'font: {fn}, {fs}')
                    break
                break
            print(f'    cell[{c_idx}]: {" | ".join(cell_info)}')
    print()
