"""Deep-dive: extract paragraph-level formatting from table cells in FYP_Report_111.docx"""
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

doc = Document(r'c:\Users\arbab\Desktop\myfrontend\FYP_Report_111.docx')

# We know tables 2-54 use the same bordered style. Analyze table 2 deeply.
for t_idx in [1, 2]:   # 0-based = table 2 and 3
    table = doc.tables[t_idx]
    print(f'\n========================================')
    print(f'TABLE {t_idx+1}: style={table.style.name}, {len(table.rows)} rows x {len(table.columns)} cols')
    print(f'========================================')

    for r_idx, row in enumerate(table.rows[:5]):   # first 5 rows
        print(f'\n  ROW {r_idx}:')
        for c_idx, cell in enumerate(row.cells[:4]):
            print(f'    CELL[{c_idx}]: "{cell.text[:60].strip()}"')
            for p_idx, para in enumerate(cell.paragraphs):
                pf = para.paragraph_format
                align = str(para.alignment)
                sb = pf.space_before
                sa = pf.space_after
                print(f'      Para[{p_idx}] align={align}, space_before={sb}, space_after={sa}')
                for run in para.runs:
                    fn   = run.font.name
                    fs   = run.font.size
                    bold = run.bold
                    ital = run.italic
                    try:
                        fc = str(run.font.color.rgb) if run.font.color and run.font.color.type else 'inherited'
                    except Exception:
                        fc = 'err'
                    print(f'      Run: name={fn}, size={fs}, bold={bold}, italic={ital}, color={fc}')

# Also check the style itself for inherited formatting
print('\n\n=== TABLE GRID STYLE DEFAULTS ===')
tg_style = doc.styles['Table Grid']
print(f'  Font name: {tg_style.font.name}')
print(f'  Font size: {tg_style.font.size}')
print(f'  Bold:      {tg_style.font.bold}')

# Check what paragraph style is used inside table cells
print('\n=== PARAGRAPH STYLES INSIDE TABLE CELLS ===')
for t_idx in [1]:
    table = doc.tables[t_idx]
    for row in table.rows[:3]:
        for cell in row.cells[:2]:
            for para in cell.paragraphs:
                print(f'  Para style: "{para.style.name}" | text: "{para.text[:40]}"')
